"""
El catálogo en hoja de cálculo y en PDF.

Tres cosas y nada más:

    a_excel(filas)   el catálogo en .xlsx, listo para abrir y editar
    de_excel(bytes)  lo contrario: lee un .xlsx y devuelve filas
    a_pdf(filas)     una lista impresa, para mirar o mandar a alguien

Qué va y qué no va en estos archivos
------------------------------------
Las FOTOS no viajan aquí. En la columna `img` va la ruta
(`/api/v1/fotos/un-dulce?v=1a2b3c`), no la imagen. Son dos razones: una hoja
de cálculo con cuatro mil caracteres de base64 por fila es inmanejable, y al
reimportar no hace falta —el importador reconoce el producto por su `id` y le
vuelve a colgar la foto que ya tenía—. O sea: se puede exportar a Excel,
cambiar precios, reimportar, y las fotos siguen ahí.

El JSON de siempre sigue siendo el respaldo completo: es el único que lleva
las fotos dentro.
"""

import io
import unicodedata
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException

# Las columnas, en orden, y de dónde sale cada una. La CABECERA es lo que lee
# `de_excel` para saber qué columna es cuál, así que el usuario puede mover las
# columnas de sitio o borrar las que no le interesen y seguirá funcionando.
COLUMNAS = [
    ("id", "Identificador", 22),
    ("nombre", "Nombre", 30),
    ("cat", "Categoría", 14),
    ("emoji", "Emoji", 8),
    ("origen", "Origen", 18),
    ("desc", "Descripción", 52),
    ("precio", "Precio", 10),
    ("etiqueta", "Etiqueta", 14),
    ("tipo", "Color etiqueta", 14),
    ("c1", "Color 1", 11),
    ("c2", "Color 2", 11),
    ("img", "Foto (ruta)", 34),
]

# Para encontrar la columna aunque venga con acentos cambiados, en mayúsculas
# o con espacios de más: «descripcion», «DESCRIPCIÓN» y «Descripción » son la
# misma. Si no, un archivo que pasó por Excel y por Google Sheets deja de
# importarse por una tilde.
def _clave(texto: Any) -> str:
    bruto = str(texto or "").strip().lower()
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", bruto) if unicodedata.category(c) != "Mn"
    )
    return " ".join(sin_tildes.split())


CABECERAS = {_clave(titulo): campo for campo, titulo, _ in COLUMNAS}
# Los nombres internos también valen como cabecera: así un .xlsx generado por
# otro sitio (o el propio JSON pasado a hoja) entra sin renombrar nada.
CABECERAS.update({_clave(campo): campo for campo, _, _ in COLUMNAS})


# ── Excel: escribir ────────────────────────────────────────────────────────

def a_excel(filas: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Catálogo"

    rosa = PatternFill("solid", fgColor="F42A8F")
    titulo = Font(bold=True, color="FFFFFF", size=11)

    for i, (_, cabecera, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=cabecera)
        celda.fill = rosa
        celda.font = titulo
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    for f, fila in enumerate(filas, start=2):
        for c, (campo, _, _) in enumerate(COLUMNAS, start=1):
            valor = fila.get(campo, "")
            if campo == "precio":
                # Número de verdad, no texto: si no, en Excel no se puede
                # sumar ni ordenar y sale con el triangulito verde de aviso.
                try:
                    valor = float(valor)
                except (TypeError, ValueError):
                    valor = 0.0
            hoja.cell(row=f, column=c, value=valor)
        hoja.cell(row=f, column=7).number_format = '"$"#,##0.00'

    # La fila de cabeceras se queda fija y con filtros: con doscientos dulces,
    # bajar y no saber qué columna se está mirando es el problema de siempre.
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(len(filas) + 1, 1)}"

    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


# ── Excel: leer ────────────────────────────────────────────────────────────

def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    return " ".join(str(valor).split()) if not isinstance(valor, str) else str(valor).strip()


def de_excel(datos: bytes) -> list[dict]:
    """
    Lee un .xlsx y devuelve las filas como diccionarios.

    Sólo mira la primera hoja y la primera fila como cabecera. Las columnas
    que no reconoce las ignora en silencio —una hoja con notas propias al lado
    tiene que poder importarse igual—, pero si no reconoce NINGUNA, para: eso
    no es un catálogo y seguir sería borrar la tienda con una lista vacía.
    """
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="No pude abrir ese archivo. Tiene que ser una hoja de cálculo .xlsx.",
        )

    try:
        hoja = libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)

        try:
            cabecera = next(filas)
        except StopIteration:
            raise HTTPException(status_code=400, detail="La hoja está vacía.")

        posiciones = {}
        for i, titulo in enumerate(cabecera or ()):
            campo = CABECERAS.get(_clave(titulo))
            if campo and campo not in posiciones:
                posiciones[campo] = i

        if not posiciones:
            esperadas = ", ".join(t for _, t, _ in COLUMNAS[:4])
            raise HTTPException(
                status_code=400,
                detail=(
                    "La primera fila de la hoja tiene que ser la cabecera con los nombres "
                    f"de las columnas ({esperadas}…). Exporta el catálogo a Excel para ver "
                    "el formato que espera."
                ),
            )
        if "nombre" not in posiciones:
            raise HTTPException(
                status_code=400,
                detail="Falta la columna «Nombre», que es la única imprescindible.",
            )

        salida: list[dict] = []
        for numero, fila in enumerate(filas, start=2):
            if fila is None or all(c is None or _texto(c) == "" for c in fila):
                continue      # fila en blanco: en una hoja a mano sobran siempre

            def dame(campo: str) -> Any:
                i = posiciones.get(campo)
                return fila[i] if i is not None and i < len(fila) else None

            nombre = _texto(dame("nombre"))
            if not nombre:
                raise HTTPException(
                    status_code=400,
                    detail=f"La fila {numero} no tiene nombre. Ponle uno o bórrala.",
                )

            salida.append(
                {
                    "id": _texto(dame("id")),
                    "nombre": nombre,
                    "cat": _texto(dame("cat")),
                    "emoji": _texto(dame("emoji")) or "🍬",
                    "origen": _texto(dame("origen")),
                    "desc": _texto(dame("desc")),
                    "precio": _precio(dame("precio"), numero),
                    "etiqueta": _texto(dame("etiqueta")),
                    "tipo": _texto(dame("tipo")),
                    "c1": _texto(dame("c1")) or "#F42A8F",
                    "c2": _texto(dame("c2")) or "#FFB8DC",
                    "img": _texto(dame("img")),
                }
            )

        if not salida:
            raise HTTPException(status_code=400, detail="La hoja no tiene ninguna fila con datos.")
        return salida
    finally:
        libro.close()


def _precio(valor: Any, fila: int) -> float:
    """
    El precio tal y como lo escribe la gente: 89, 89.5, «$89.50», «1,299».

    Excel suele dar un número ya, pero en cuanto alguien pega de otro sitio
    llega como texto con el símbolo delante.
    """
    if valor is None or _texto(valor) == "":
        return 0.0
    if isinstance(valor, (int, float, Decimal)):
        return float(valor)

    limpio = _texto(valor).replace("$", "").replace(",", "").replace(" ", "")
    try:
        return float(Decimal(limpio))
    except (InvalidOperation, ValueError):
        raise HTTPException(
            status_code=400,
            detail=f"El precio de la fila {fila} («{_texto(valor)}») no es un número.",
        )


# ── PDF ────────────────────────────────────────────────────────────────────

# Las fuentes que trae ReportLab de serie llegan hasta Latin-1: los acentos y
# la ñ entran, los emojis no. Antes que incrustar una fuente de un mega en la
# imagen para dibujar un 🍫 en una tabla, se quitan y ya.
def _imprimible(texto: Any) -> str:
    bruto = str(texto or "")
    return "".join(c for c in bruto if ord(c) < 256).strip()


def a_pdf(filas: list[dict], titulo: str = "Catálogo") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    memoria = io.BytesIO()
    documento = SimpleDocTemplate(
        memoria,
        pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=14 * mm,
        title=_imprimible(titulo), author="Candylandia Store",
    )

    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("titulo", parent=estilos["Title"], fontSize=18,
                        textColor=colors.HexColor("#6E2A8C"), spaceAfter=2)
    pie = ParagraphStyle("pie", parent=estilos["Normal"], fontSize=8,
                         textColor=colors.HexColor("#7A6B80"))
    celda = ParagraphStyle("celda", parent=estilos["Normal"], fontSize=8, leading=10)
    celda_der = ParagraphStyle("celdaDer", parent=celda, alignment=TA_RIGHT)
    cabecera = ParagraphStyle("cab", parent=celda, fontSize=8.5,
                              textColor=colors.white, fontName="Helvetica-Bold")

    # Sin la columna de emoji ni la de colores: en papel no aportan nada.
    campos = [
        ("nombre", "Producto", 46),
        ("cat", "Categoría", 24),
        ("origen", "Origen", 26),
        ("desc", "Descripción", 95),
        ("etiqueta", "Etiqueta", 24),
        ("precio", "Precio", 20),
    ]

    datos = [[Paragraph(t, cabecera) for _, t, _ in campos]]
    for fila in filas:
        linea = []
        for campo, _, _ in campos:
            if campo == "precio":
                try:
                    texto = f"$ {float(fila.get('precio') or 0):,.2f}"
                except (TypeError, ValueError):
                    texto = "$ 0.00"
                linea.append(Paragraph(texto, celda_der))
            else:
                linea.append(Paragraph(_imprimible(fila.get(campo, "")), celda))
        datos.append(linea)

    tabla = Table(
        datos,
        colWidths=[a * mm for _, _, a in campos],
        repeatRows=1,          # la cabecera se repite en cada hoja
    )
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F42A8F")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E8DCEF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FDF4FA")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))

    total = sum(float(f.get("precio") or 0) for f in filas)
    historia = [
        Paragraph(_imprimible(titulo), h1),
        Paragraph(
            f"{len(filas)} producto(s) &nbsp;·&nbsp; suma de precios: $ {total:,.2f}",
            pie,
        ),
        Spacer(1, 7),
        tabla,
    ]

    def _numerar(lienzo, _doc):
        """El número de página, abajo a la derecha."""
        lienzo.saveState()
        lienzo.setFont("Helvetica", 7.5)
        lienzo.setFillColor(colors.HexColor("#7A6B80"))
        ancho, _alto = landscape(A4)
        lienzo.drawRightString(ancho - 12 * mm, 8 * mm, f"Pagina {lienzo.getPageNumber()}")
        lienzo.restoreState()

    documento.build(historia, onFirstPage=_numerar, onLaterPages=_numerar)
    return memoria.getvalue()
